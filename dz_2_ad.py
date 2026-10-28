import requests
from bs4 import BeautifulSoup
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt
from datetime import datetime
import re

# --- НАСТРОЙКИ ---
URL = "https://www.mirea.ru/news/index.php?set_filter=Y&arrFilter_ff[TAGS]=абитуриентам"
CSV_FILE = "mirea_news.csv"
GRAPH_YEAR = "news_by_year.png"
GRAPH_MONTH = "news_by_month.png"
DOCX_FILE = "mirea_news_report.docx"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"
}

def parse_page(html):
    """Извлекает новости только с одной страницы."""
    soup = BeautifulSoup(html, "html.parser")
    items = []

    # Стратегия 1: ищем карточки новостей по классу .news-item (частый шаблон Битрикса)
    for card in soup.select(".news-item"):
        title_tag = card.select_one("h2 a, .news-title a, h3 a")
        date_tag = card.select_one(".news-date, .date, span.date")

        if title_tag and date_tag:
            title = title_tag.get_text(strip=True)
            date_str = date_tag.get_text(strip=True)
            items.append((title, date_str))

    # Стратегия 2: если карточек нет, ищем статьи через <article>
    if not items:
        for article in soup.select("article"):
            title_tag = article.select_one("h2, h3, .title")
            date_tag = article.select_one(".date, time")
            if title_tag and date_tag:
                title = title_tag.get_text(strip=True)
                date_str = date_tag.get_text(strip=True)
                items.append((title, date_str))

    return items

def normalize_date(date_str):
    """Приводит дату к формату DD.MM.YYYY, если возможно."""
    # Ожидаем форматы вроде '15.10.2024', '15 октября 2024' и т.п.
    # Сначала пробуем прямой разбор DD.MM.YYYY
    try:
        dt = datetime.strptime(date_str, "%d.%m.%Y")
        return dt.strftime("%d.%m.%Y")
    except ValueError:
        pass

    # Если не вышло — пробуем разобрать словами (упрощённо)
    months_ru = {
        "января": "01", "февраля": "02", "марта": "03", "апреля": "04",
        "мая": "05", "июня": "06", "июля": "07", "августа": "08",
        "сентября": "09", "октября": "10", "ноября": "11", "декабря": "12"
    }
    match = re.match(r"(\d{1,2})\s+([а-я]+)\s+(\d{4})", date_str.lower())
    if match:
        day, month_name, year = match.groups()
        month = months_ru.get(month_name)
        if month:
            return f"{day.zfill(2)}.{month}.{year}"

    # Если ничего не подошло — возвращаем как есть (будет обработано позже)
    return date_str

def main():
    print("Загружаем первую страницу...")
    resp = requests.get(URL, headers=HEADERS, timeout=15)
    resp.raise_for_status()

    items = parse_page(resp.text)
    if not items:
        print("Новости не найдены. Проверьте HTML-структуру страницы в браузере.")
        return

    # Нормализуем даты
    data = []
    for title, date_raw in items:
        date_norm = normalize_date(date_raw)
        data.append({"Заголовок": title, "Дата": date_norm})

    df = pd.DataFrame(data)
    # Сохраняем в CSV с BOM для корректного открытия в Excel
    df.to_csv(CSV_FILE, index=False, encoding="utf-8-sig")
    print(f"Сохранено {len(df)} новостей в {CSV_FILE}")

    # Приводим даты к datetime
    df["Дата_dt"] = pd.to_datetime(df["Дата"], format="%d.%m.%Y", errors="coerce")
    df_clean = df.dropna(subset=["Дата_dt"])

    if df_clean.empty:
        print("Не удалось распознать ни одной даты. Пропускаем графики и отчёт.")
        return

    # Распределение по годам
    df_clean["Год"] = df_clean["Дата_dt"].dt.year
    year_counts = df_clean.groupby("Год").size().reset_index(name="Количество")
    year_counts = year_counts.sort_values("Год")

    plt.figure(figsize=(10, 6))
    plt.bar(year_counts["Год"].astype(str), year_counts["Количество"], color="#4a90e2")
    plt.title("Количество новостей по годам", fontsize=14, fontweight="bold")
    plt.xlabel("Год", fontsize=12)
    plt.ylabel("Количество публикаций", fontsize=12)
    plt.grid(axis="y", linestyle="--", alpha=0.3)
    plt.tight_layout()
    plt.savefig(GRAPH_YEAR, dpi=300)
    plt.close()
    print(f"График по годам сохранён: {GRAPH_YEAR}")

    # Распределение по месяцам
    df_clean["Месяц"] = df_clean["Дата_dt"].dt.to_period("M").astype(str)
    month_counts = df_clean.groupby("Месяц").size().reset_index(name="Количество")
    month_counts["Месяц_dt"] = pd.to_datetime(month_counts["Месяц"], format="%Y-%m")
    month_counts = month_counts.sort_values("Месяц_dt")

    plt.figure(figsize=(10, 6))
    plt.bar(month_counts["Месяц"].astype(str), month_counts["Количество"], color="#2ecc71")
    plt.title("Количество новостей по месяцам", fontsize=14, fontweight="bold")
    plt.xlabel("Месяц", fontsize=12)
    plt.ylabel("Количество публикаций", fontsize=12)
    plt.xticks(rotation=45)
    plt.grid(axis="y", linestyle="--", alpha=0.3)
    plt.tight_layout()
    plt.savefig(GRAPH_MONTH, dpi=300)
    plt.close()
    print(f"График по месяцам сохранён: {GRAPH_MONTH}")

    # Генерация DOCX-отчёта
    doc = Document()
    doc.add_heading("Отчёт по новостям для абитуриентов МИРЭА (первая страница)", level=1)
    doc.add_paragraph(
        f"Отчёт сформирован на основе первой страницы выдачи по фильтру «абитуриентам». "
        f"Всего обработано {len(df_clean)} публикаций с распознанными датами."
    )

    doc.add_heading("Распределение публикаций по месяцам", level=2)
    doc.add_picture(GRAPH_MONTH, width=Inches(6))

    doc.add_heading("Список публикаций", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "№"
    hdr_cells[1].text = "Дата"
    hdr_cells[2].text = "Заголовок"
    for cell in table.rows[0].cells:
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(11)

    for idx, row in df_clean.iterrows():
        r_cells = table.add_row().cells
        r_cells[0].text = str(idx + 1)
        r_cells[1].text = row["Дата_dt"].strftime("%d.%m.%Y")
        r_cells[2].text = row["Заголовок"]

    doc.save(DOCX_FILE)
    print(f"Отчёт сохранён: {DOCX_FILE}")

if __name__ == "__main__":
    main()