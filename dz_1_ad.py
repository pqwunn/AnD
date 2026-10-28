#Парсер новостей МАИ

import re
import time
import requests
from bs4 import BeautifulSoup
from datetime import datetime

import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from docx import Document
from docx.shared import Inches


BASE_URL = "https://mai.ru/press/news/?tags=7462"
HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) '
                  'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8',
    'Accept-Language': 'ru-RU,ru;q=0.9,en;q=0.8',
}

MAX_PAGES = 50
DELAY = 0.5

MONTHS_RU = {
    'янв': 1, 'фев': 2, 'мар': 3, 'апр': 4,
    'май': 5, 'мая': 5, 'июн': 6, 'июл': 7,
    'авг': 8, 'сен': 9, 'окт': 10, 'ноя': 11, 'дек': 12,
}

def parse_mai_date(date_text, default_year=None):
    if default_year is None:
        default_year = datetime.now().year

    text = date_text.strip().lower()
    m = re.match(r'(\d{1,2})\s+([а-яё]+)(?:\s+(\d{4}))?', text)
    if not m:
        return None

    day = int(m.group(1))
    month_key = m.group(2)[:3]
    year_str = m.group(3)

    month = MONTHS_RU.get(month_key)
    if not month:
        return None

    year = int(year_str) if year_str else default_year

    try:
        return datetime(year, month, day)
    except ValueError:
        return None

def parse_page(html, default_year=None):
    soup = BeautifulSoup(html, 'lxml')
    news_list = []
    if default_year is None:
        default_year = datetime.now().year

    articles = soup.find_all('article')

    for art in articles:
        h5 = art.find('h5', class_='mb-3')
        if not h5:
            continue
        title = h5.get_text(strip=True)

        date_span = art.find('span', class_='badge')
        if not date_span:
            continue
        date_text = date_span.get_text(strip=True)

        date_obj = parse_mai_date(date_text, default_year)
        if not date_obj:
            print(f" Не распознана дата: {date_text!r}")
            continue

        a_tag = art.find_parent('a')
        href = a_tag.get('href', '') if a_tag else ''
        if href.startswith('http'):
            url = href
        else:
            url = f"https://mai.ru/press/news/{href}" if href else ''

        news_list.append({
            'title': title,
            'url': url,
            'date': date_obj,
            'year': date_obj.year,
            'month': date_obj.month,
        })

    return news_list

def find_last_page(html):
    soup = BeautifulSoup(html, 'lxml')
    max_page = 1

    for a in soup.find_all('a', href=True):
        href = a['href']
        m = re.search(r'PAGEN_1=(\d+)', href) or re.search(r'[?&]page=(\d+)', href)
        if m:
            max_page = max(max_page, int(m.group(1)))

    return max_page

def parse_all_pages():
    session = requests.Session()
    session.headers.update(HEADERS)

    try:
        session.get("https://mai.ru/", timeout=20)
    except Exception as e:
        print(f"Не удалось получить cookies: {e}")

    all_news = []
    seen = set()

    print(f"[1] {BASE_URL}")
    try:
        r = session.get(BASE_URL, timeout=30)
        r.raise_for_status()
    except Exception as e:
        print(f" Ошибка на первой странице: {e}")
        print("Попробуй Selenium")
        return []

    html_first = r.text
    total_pages = find_last_page(html_first)
    print(f"Найдено страниц: {total_pages}")

    with open('mai_debug.html', 'w', encoding='utf-8') as f:
        f.write(html_first)

    page_news = parse_page(html_first)
    print(f"новостей на странице: {len(page_news)}")
    for n in page_news:
        key = (n['title'], n['date'])
        if key not in seen:
            seen.add(key)
            all_news.append(n)

    for page in range(2, min(total_pages, MAX_PAGES) + 1):
        if '?' in BASE_URL:
            url = f"{BASE_URL}&PAGEN_1={page}"
        else:
            url = f"{BASE_URL}?PAGEN_1={page}"

        print(f"[{page}/{total_pages}] {url}")
        try:
            r = session.get(url, timeout=30)
            r.raise_for_status()
        except Exception as e:
            print(f"  ⚠ Ошибка: {e}")
            continue

        page_news = parse_page(r.text)
        print(f"новостей: {len(page_news)}")

        if not page_news:
            print("  Пустая страница — останавливаемся.")
            break

        for n in page_news:
            key = (n['title'], n['date'])
            if key not in seen:
                seen.add(key)
                all_news.append(n)

        time.sleep(DELAY)

    print(f"\n✅ Всего собрано новостей: {len(all_news)}")
    return all_news

def save_to_csv(data, filename='mai_news.csv'):
    if not data:
        print("Нет данных для CSV.")
        return
    df = pd.DataFrame(data).sort_values('date', ascending=False)
    df.to_csv(filename, index=False, encoding='utf-8-sig')
    print(f"💾 CSV: {filename}")

def plot_distribution(data, filename='mai_news_chart.png'):
    if not data:
        print("Нет данных для графика.")
        return None

    df = pd.DataFrame(data)
    df['period'] = df['date'].dt.to_period('M')
    counts = df.groupby('period').size().sort_index()

    plt.figure(figsize=(14, 6))
    counts.plot(kind='bar', color='steelblue', edgecolor='black')
    plt.title('Распределение новостей МАИ (абитуриентам) по месяцам')
    plt.xlabel('Месяц')
    plt.ylabel('Количество новостей')
    plt.xticks(rotation=60, ha='right')
    plt.grid(axis='y', linestyle='--', alpha=0.6)
    plt.tight_layout()
    plt.savefig(filename, dpi=150)
    plt.close()
    print(f"📊 График: {filename}")
    return filename

def create_report(data, chart_path, filename='mai_news_report.docx'):
    if not data:
        print("Нет данных для отчёта.")
        return

    doc = Document()
    doc.add_heading('Отчёт по новостям МАИ для абитуриентов', 0)
    doc.add_paragraph(f'Дата генерации: {datetime.now():%d.%m.%Y %H:%M}')
    doc.add_paragraph(f'Всего новостей: {len(data)}')

    if chart_path:
        doc.add_heading('Распределение по месяцам', level=1)
        doc.add_picture(chart_path, width=Inches(6.5))

    doc.add_heading('Список публикаций', level=1)
    df = pd.DataFrame(data).sort_values('date', ascending=False)

    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text = 'Дата'
    hdr[1].text = 'Заголовок'
    hdr[2].text = 'Ссылка'

    for _, row in df.iterrows():
        cells = table.add_row().cells
        cells[0].text = row['date'].strftime('%d.%m.%Y')
        cells[1].text = row['title']
        cells[2].text = row['url']

    doc.save(filename)
    print(f"📄 DOCX: {filename}")

if __name__ == "__main__":
    print("Парсинг новостей МАИ")
    news = parse_all_pages()

    if news:
        save_to_csv(news)
        chart = plot_distribution(news)
        create_report(news, chart)
        print("\n✅ Готово")
    else:
        print("\n❌ Не удалось собрать данные.")